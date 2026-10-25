import openpyxl
import re
import sys
import os, sys, tempfile, unittest, subprocess


def _read(path):
    with open(path, encoding='utf-8') as fh:
        return fh.read()

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import mapper


class T(unittest.TestCase):
    def test_dates(self):
        w = []
        self.assertEqual(mapper.parse_date('1980-05-02')[0], '2 MAY 1980')
        self.assertEqual(mapper.parse_date('c. 1855-1856')[0], 'BET 1855 AND 1856')
        self.assertEqual(mapper.parse_date('c. 1890'), ('ABT 1890', 'c. 1890'))
        self.assertEqual(mapper.parse_date('1952')[0], '1952')
        self.assertEqual(mapper.parse_date('1900-00-10', w), (None, '1900-00-10'))
        self.assertEqual(mapper.parse_date('1900-02-31', w), (None, '1900-02-31'))
        self.assertEqual(mapper.parse_date('1900-13', w), (None, '1900-13'))
        self.assertEqual(mapper.parse_date('c. 1980-1970', w), (None, 'c. 1980-1970'))
        self.assertEqual(len(w), 4)

    def test_no_truncation(self):
        txt = 'x' * 700 + '\nsegunda'
        lines = mapper.gedcom_text(1, 'NOTE', txt)
        rebuilt = ''
        for ln in lines:
            tag = ln.split(' ', 2)
            body = tag[2] if len(tag) > 2 else ''
            if tag[1] == 'CONT':
                rebuilt += '\n' + body
            else:
                rebuilt += body
        self.assertEqual(rebuilt, txt)
        self.assertTrue(all(len(l) <= 250 for l in lines))

    def test_fixture_runs_and_dedup(self):
        d = tempfile.mkdtemp()
        subprocess.check_call([sys.executable, os.path.join(os.path.dirname(mapper.__file__), 'make_fixture.py')], cwd=d)
        s, warn, rep, unsup = mapper.run(os.path.join(d, 'fixtures/fixture_sintetico.xlsx'), os.path.join(d, 'o.ged'))
        self.assertEqual(s['personas'], 3)
        self.assertEqual(s['familias'], 1)  # pareja + hijo = una sola familia
        self.assertEqual(s['citas_mapeadas'], 1)
        self.assertEqual(s['citas_excluidas_confianza'], 1)
        self.assertEqual(unsup, [])

    def _fixture(self):
        import openpyxl
        d = tempfile.mkdtemp()
        subprocess.check_call([sys.executable, os.path.join(os.path.dirname(mapper.__file__), 'make_fixture.py')], cwd=d)
        p = os.path.join(d, 'fixtures/fixture_sintetico.xlsx')
        return d, p, openpyxl.load_workbook(p)

    def test_no_duplicate_family_when_sex_missing(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        for r in (3, 4):
            ws.cell(r, 5).value = ''          # sexo vacio en la pareja
        ws.cell(2, 10).value = ''; ws.cell(2, 11).value = ''   # sin hijos: solo familia por conyuge
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertEqual(s['familias'], 1)
        self.assertTrue(any('sexo no determinante' in w for w in warn))

    def test_duplicate_ids_fail_closed(self):
        d, p, wb = self._fixture()
        wb['30_Arbol_Confirmado'].cell(3, 1).value = 'P000001'
        wb.save(p)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'))

    def test_missing_source_is_excluded_with_note(self):
        d, p, wb = self._fixture()
        wb['05_Citas'].cell(5, 2).value = 'S999999'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertEqual(s['citas_excluidas_sin_fuente'], 1)
        self.assertIn('EXCLUIDA', _read(os.path.join(d, 'o.ged')))

    def _out(self, d):
        return _read(os.path.join(d, 'o.ged'))

    def test_privacy_fail_closed_and_no_filenames(self):
        d, p, wb = self._fixture()
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertGreaterEqual(s['personas_privadas_RESN'], 1)
        self.assertIn('RESN privacy', o)
        self.assertNotIn('doc_ficticio.pdf', o)          # privacidad desconocida/Privado: sin nombre de archivo
        self.assertIn('archivo omitido por privacidad', o)

    def test_deceased_old_person_with_public_doc_is_not_private(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        ws.cell(2, 6).value = '1850-01-01'; ws.cell(2, 8).value = '1920-01-01'
        d6 = wb['06_Documentos']; d6.cell(5, 13).value = 'Publico'
        wb.save(p)
        mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertIn('doc_ficticio.pdf', o)

    def test_no_none_literal(self):
        d, p, wb = self._fixture()
        mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertNotIn('None', self._out(d))

    def test_citation_url_preserved_in_note(self):
        d, p, wb = self._fixture()
        wb['05_Citas'].cell(5, 12).value = 'https://ejemplo.invalid/ark:/1/' + 'a' * 300
        ws = wb['30_Arbol_Confirmado']; ws.cell(2, 6).value = '1850-01-01'; ws.cell(2, 8).value = '1920-01-01'   # persona publica
        wb['06_Documentos'].cell(5, 13).value = 'Público'
        wb.save(p)
        mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d).replace('\n3 CONC ', '').replace('\n4 CONC ', '')
        self.assertIn('https://ejemplo.invalid/ark:/1/' + 'a' * 300, o)

    def _priv_fixture(self, public):
        """P000001 (hijo, C000001) privado o publico; datos marcadores en nota personal, cita y documento."""
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        ws.cell(2, 17).value = 'NOTA_PERSONAL_SECRETA'
        ws.cell(3, 17).value = 'NOTA_OTRA_PERSONA'
        c = wb['05_Citas']
        c.cell(5, 5).value = 'VALOR_SECRETO'; c.cell(5, 8).value = 'EVAL_SECRETA'; c.cell(5, 11).value = 'D000001'; c.cell(5, 12).value = 'https://x.invalid/URL_SECRETA'
        c.cell(5, 6).value = 'Acte n°12, vue 34'
        if public:
            ws.cell(2, 6).value = '1850-01-01'; ws.cell(2, 8).value = '1920-01-01'
            wb['06_Documentos'].cell(5, 13).value = 'Público'
        wb.save(p)
        return d, p

    def test_private_person_omits_personal_note_and_sensitive_citation_fields(self):
        d, p = self._priv_fixture(public=False)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        for marcador in ('NOTA_PERSONAL_SECRETA', 'VALOR_SECRETO', 'EVAL_SECRETA', 'URL_SECRETA'):
            self.assertNotIn(marcador, o, marcador)
        self.assertIn('Nota personal omitida por privacidad (ver Excel)', o)
        self.assertIn('PAGE Acte n°12, vue 34 | C000001 | Nacimiento\n', o)       # localizador seguro, ID y tipo se conservan
        self.assertIn('QUAY 2', o)
        self.assertIn('Corroborada: Sí', o)
        self.assertIn('Valor afirmado, Evaluacion, Doc, URL y Extracto omitidos por privacidad', o)
        self.assertEqual(s['notas_personales_omitidas_privacidad'], len(s['omisiones_privacidad_ids']['nota_personal']))
        self.assertIn('P000001', s['omisiones_privacidad_ids']['nota_personal'])
        self.assertIn('C000001', s['omisiones_privacidad_ids']['cita_campos'])
        self.assertEqual(s['citas_con_campos_omitidos_privacidad'], len(s['omisiones_privacidad_ids']['cita_campos']))

    def test_public_person_keeps_note_and_citation_fields(self):
        d, p = self._priv_fixture(public=True)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertIn('NOTA_PERSONAL_SECRETA', o)
        for marcador in ('VALOR_SECRETO', 'EVAL_SECRETA', 'URL_SECRETA'):
            self.assertIn(marcador, o, marcador)
        self.assertNotIn('P000001', s['omisiones_privacidad_ids']['nota_personal'])
        self.assertNotIn('C000001', s['omisiones_privacidad_ids']['cita_campos'])

    def test_private_person_unsafe_locator_is_omitted_with_reason_counter(self):
        for loc in ('DNI 12345678Z', 'ref 123456789', 'http://x.invalid/a', 'mail a@b.c', 'x' * 130):
            d, p = self._priv_fixture(public=False)
            import openpyxl
            wb = openpyxl.load_workbook(p); wb['05_Citas'].cell(5, 6).value = loc; wb.save(p)
            s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
            o = self._out(d)
            self.assertIn('PAGE localizador omitido por privacidad | C000001 | Nacimiento', o, loc)
            self.assertNotIn(loc, o, loc)
            self.assertEqual(s['omisiones_privacidad_ids']['localizador'], ['C000001'], loc)
            self.assertEqual(s['localizadores_omitidos_privacidad'], 1)

    def test_private_marr_citation_omits_fields_too(self):
        d, p = self._marr_setup('P000002', 'P000003')
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertNotIn('EXTRACTO_SECRETO', o)
        self.assertIn('C000001', s['omisiones_privacidad_ids']['cita_campos'])

    def _src_fixture(self, titulo='Registro parroquial 1850', autor='Parroquia X', sig='Libro 3', doc='Público'):
        d, p, wb = self._fixture()
        f = wb['04_Fuentes']
        f.cell(5, 3).value = titulo; f.cell(5, 4).value = autor; f.cell(5, 7).value = sig
        wb['06_Documentos'].cell(5, 13).value = doc
        wb['06_Documentos'].cell(5, 8).value = 'S000001'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        return s, self._out(d)

    def test_source_old_public_keeps_title_author_and_signatura(self):
        s, o = self._src_fixture()
        self.assertIn('1 TITL Registro parroquial 1850', o); self.assertIn('1 AUTH Parroquia X', o); self.assertIn('Signatura: Libro 3', o)
        self.assertEqual(s['omisiones_privacidad_ids']['fuente_restringida'], [])

    def test_source_with_nonpublic_document_is_restricted_with_reason(self):
        s, o = self._src_fixture(titulo='TITULO_SECRETO_PERSONA', autor='AUTOR_SECRETO', doc='Privado')
        self.assertNotIn('TITULO_SECRETO_PERSONA', o); self.assertNotIn('AUTOR_SECRETO', o)
        self.assertIn('1 TITL Fuente restringida S000001 (ver Excel)', o)
        self.assertEqual(s['omisiones_privacidad_ids']['fuente_restringida'], ['S000001'])
        self.assertIn('S000001: documento no publico', s['omisiones_privacidad_ids']['fuente_restringida_motivos'])
        self.assertEqual(s['fuentes_restringidas_privacidad'], 1)

    def test_source_identity_document_title_is_restricted_even_if_doc_public(self):
        s, o = self._src_fixture(titulo='DNI de Persona Inventada')
        self.assertNotIn('Persona Inventada', o)
        self.assertIn('fuente de documento de identidad', ' '.join(s['omisiones_privacidad_ids']['fuente_restringida_motivos']))

    def test_source_personal_testimony_or_mail_is_restricted(self):
        for tipo, repo in (('Testimonio oral/directo', 'Proyecto'), ('Registro', 'Conversacion WhatsApp con X'), ('Correo oficial', 'Gmail')):
            d, p, wb = self._fixture()
            f = wb['04_Fuentes']; f.cell(5, 2).value = tipo; f.cell(5, 3).value = 'TITULO_PERSONAL'; f.cell(5, 5).value = repo
            wb.save(p); s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
            self.assertNotIn('TITULO_PERSONAL', self._out(d), tipo)
            self.assertIn('testimonio, mensaje o correo personal', ' '.join(s['omisiones_privacidad_ids']['fuente_restringida_motivos']))

    def test_document_note_for_private_person_omits_type_and_privacy_text(self):
        d, p, wb = self._fixture()
        dd = wb['06_Documentos']; dd.cell(5, 3).value = 'DNI de Maria Viva'; dd.cell(5, 13).value = 'Privado (DNI 12345678A)'; dd.cell(5, 14).value = ''
        wb.save(p); s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        for m in ('Maria Viva', '12345678A', 'doc_ficticio'):
            self.assertNotIn(m, o, m)
        self.assertIn('Documento D000001: archivo omitido por privacidad | tipo omitido por privacidad (ver Excel)', o)
        self.assertIn('privacidad omitida por privacidad (ver Excel)', o)
        self.assertEqual(s['omisiones_privacidad_ids']['documento_campos'], ['D000001'])
        self.assertEqual(s['documentos_con_campos_omitidos_privacidad'], 1)

    def test_document_note_public_doc_of_public_person_keeps_type(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']; ws.cell(2, 6).value = '1850-01-01'; ws.cell(2, 8).value = '1920-01-01'
        dd = wb['06_Documentos']; dd.cell(5, 3).value = 'Acta de bautismo'; dd.cell(5, 13).value = 'Público'; dd.cell(5, 8).value = 'S000001'
        wb.save(p); s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertEqual(s['omisiones_privacidad_ids']['documento_campos'], [])
        self.assertIn('tipo Acta de bautismo', self._out(d))

    def test_source_url_with_eight_digits_omitted_but_normal_url_kept(self):
        for url, omitted in (('https://archivo.example/12345678/doc', True), ('https://archivo.example/ark/abc123/doc', False)):
            d, p, wb = self._fixture()
            wb['04_Fuentes'].cell(5, 11).value = url
            wb['06_Documentos'].cell(5, 13).value = 'Público'
            wb.save(p); s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
            o = self._out(d)
            self.assertEqual(url in o, not omitted, url)
            self.assertEqual('S000001:URL' in s['omisiones_privacidad_ids']['fuente_campo'], omitted)

    def test_default_restricted_source_ids_are_restricted_with_reason(self):
        d, p, wb = self._fixture()
        wb['04_Fuentes'].cell(5, 1).value = 'S000112'
        for r in range(5, wb['05_Citas'].max_row + 1):
            if wb['05_Citas'].cell(r, 2).value == 'S000001': wb['05_Citas'].cell(r, 2).value = 'S000112'
        for r in range(5, wb['06_Documentos'].max_row + 1):
            if wb['06_Documentos'].cell(r, 8).value == 'S000001': wb['06_Documentos'].cell(r, 8).value = 'S000112'
        for r in range(2, wb['30_Arbol_Confirmado'].max_row + 1):
            v = wb['30_Arbol_Confirmado'].cell(r, 15).value
            if v and 'S000001' in str(v): wb['30_Arbol_Confirmado'].cell(r, 15).value = str(v).replace('S000001', 'S000112')
        wb.save(p); s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertIn('S000112', s['omisiones_privacidad_ids']['fuente_restringida'])
        self.assertIn('pendiente de validar', ' '.join(s['omisiones_privacidad_ids']['fuente_restringida_motivos']))
        self.assertNotIn('Registro ficticio', self._out(d))

    def test_default_omitted_locators_for_private_citations(self):
        d, p = self._priv_fixture(public=False)
        wb = openpyxl.load_workbook(p); wb['05_Citas'].cell(5, 1).value = 'C000188'; wb['05_Citas'].cell(5, 6).value = 'Correo oficial de Persona Viva, 11/09/2026'; wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertNotIn('Persona Viva', o)
        self.assertIn('PAGE localizador omitido por privacidad | C000188 | Nacimiento', o)
        self.assertEqual(s['omisiones_privacidad_ids']['localizador'], ['C000188'])

    def test_source_with_recent_year_in_title_is_restricted(self):
        s, o = self._src_fixture(titulo='Acta de defuncion de Fulano, 2009')
        self.assertNotIn('Fulano', o)
        self.assertEqual(s['omisiones_privacidad_ids']['fuente_restringida'], ['S000001'])

    def test_source_unsafe_signatura_is_omitted_and_counted(self):
        s, o = self._src_fixture(sig='ref DNI 12345678Z')
        self.assertNotIn('12345678Z', o)
        self.assertEqual(s['omisiones_privacidad_ids']['fuente_campo'], ['S000001:Signatura'])
        self.assertIn('1 TITL Registro parroquial 1850', o)

    def test_invalid_date_keeps_original_note(self):
        d, p, wb = self._fixture()
        wb['30_Arbol_Confirmado'].cell(3, 6).value = '1900-02-31'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertIn('1900-02-31', self._out(d))
        self.assertTrue(any('fecha no mapeada' in w for w in warn))

    def test_conflicting_marriage_dates(self):
        d, p, wb = self._fixture()
        wb['30_Arbol_Confirmado'].cell(4, 13).value = '1976-02-02'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertNotIn('1 MARR', o)
        self.assertIn('CONTRADICTORIOS', o)
        self.assertTrue(any('contradictorios' in w for w in warn))

    def test_malformed_xref_and_collision(self):
        d, p, wb = self._fixture()
        wb['30_Arbol_Confirmado'].cell(3, 1).value = 'P 000002'
        wb.save(p)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'))
        d, p, wb = self._fixture()
        wb['04_Fuentes'].cell(5, 1).value = 'P000001'
        wb.save(p)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'))

    def test_reconciliation_counts_out_of_scope(self):
        d, p, wb = self._fixture()
        wb['05_Citas'].cell(5, 3).value = 'P999999'   # persona fuera del arbol confirmado
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertEqual(s['citas_fuera_del_alcance'], 1)
        self.assertEqual(s['citas_mapeadas'] + s['citas_excluidas_confianza'] + s['citas_excluidas_sin_fuente'] + s['citas_fuera_del_alcance'], s['citas_filas_total'])

    def test_guards_missing_file_and_sheet(self):
        d, p, wb = self._fixture()
        with self.assertRaises(mapper.MapperError):
            mapper.run(os.path.join(d, 'no_existe.xlsx'), os.path.join(d, 'o.ged'))
        del wb['06_Documentos']
        wb.save(p)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'))

    def test_duplicate_person_ids_in_01_used_by_tree_fail_closed(self):
        d, p, wb = self._fixture()
        ws = wb.create_sheet('01_Personas')
        for _ in range(4): ws.append([])
        ws.append(['P000002', 'a']); ws.append(['P000002', 'b'])
        wb.save(p)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'))

    def test_duplicate_person_ids_in_01_outside_tree_only_warn(self):
        d, p, wb = self._fixture()
        ws = wb.create_sheet('01_Personas')
        for _ in range(4): ws.append([])
        ws.append(['P000900', 'a']); ws.append(['P000900', 'b'])
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertTrue(any('01_Personas con IDs duplicados' in w for w in warn))

    def test_citation_bound_to_birth_event(self):
        d, p, wb = self._fixture()
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertEqual(s['citas_ligadas_a_evento'], 1)
        self.assertIn('1 BIRT\n2 DATE 2 MAY 1980\n2 PLAC Villa Falsa\n2 SOUR @S000001@', o)

    def test_citation_with_mismatching_year_stays_on_person(self):
        d, p, wb = self._fixture()
        wb['05_Citas'].cell(5, 5).value = '1999-01-01'     # Valor_afirmado distinto del nacimiento 1980
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        self.assertEqual(s['citas_ligadas_a_evento'], 0)
        self.assertEqual(s['citas_ligadas_a_persona'], 1)
        self.assertIn('año del valor de la cita no concuerda con el evento', s['citas_persona_por_motivo'])

    def test_non_exact_tipo_hecho_stays_on_person(self):
        d, p, wb = self._fixture()
        wb['05_Citas'].cell(5, 4).value = 'Nacimiento y filiación'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        self.assertEqual(s['citas_ligadas_a_evento'], 0)
        self.assertEqual(s['citas_persona_por_motivo'].get('tipo_no_exacto'), 1)

    def test_birth_citation_without_birth_event_stays_on_person(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        ws.cell(2, 6).value = ''; ws.cell(2, 7).value = ''
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        self.assertEqual(s['citas_ligadas_a_evento'], 0)
        self.assertEqual(s['citas_persona_por_motivo'].get('evento_ausente'), 1)

    def test_marriage_citation_bound_only_with_confirmed_union(self):
        d, p, wb = self._fixture()
        c = wb['05_Citas']; c.cell(5, 3).value = 'P000002'; c.cell(5, 4).value = 'Matrimonio'; c.cell(5, 5).value = '1975-01-01'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertEqual(s['citas_ligadas_a_evento'], 1)
        self.assertIn('1 MARR\n2 DATE 1 JAN 1975\n2 PLAC Villa Falsa\n2 SOUR @S000001@', o)
        d, p, wb = self._fixture()
        r = wb['02_Relaciones']
        r.cell(7, 3).value = 'Otro'                      # quitar la fila Conyuge: union no acreditada
        c = wb['05_Citas']; c.cell(5, 3).value = 'P000002'; c.cell(5, 4).value = 'Matrimonio'; c.cell(5, 5).value = '1975-01-01'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        self.assertEqual(s['citas_ligadas_a_evento'], 0)
        self.assertEqual(s['citas_persona_por_motivo'].get('matrimonio_sin_familia_unica_confirmada'), 1)

    def _marr_setup(self, cite_person, doc_person):
        """Ambos conyuges fallecidos antiguos (publicos); el documento 'Privado' cuelga de doc_person."""
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        for r in (3, 4):
            ws.cell(r, 6).value = '1850-01-01'; ws.cell(r, 8).value = '1920-01-01'
        wb['06_Documentos'].cell(5, 4).value = doc_person
        c = wb['05_Citas']
        c.cell(5, 3).value = cite_person; c.cell(5, 4).value = 'Matrimonio'; c.cell(5, 5).value = '1975-01-01'
        c.cell(5, 7).value = 'EXTRACTO_SECRETO'
        wb.save(p)
        return d, p

    def test_marr_extract_hidden_when_spouse_not_citing_is_private(self):
        d, p = self._marr_setup('P000002', 'P000003')
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertEqual(s['citas_ligadas_a_evento'], 1)
        self.assertIn('1 MARR', o)
        self.assertNotIn('EXTRACTO_SECRETO', o)

    def test_marr_extract_hidden_other_direction(self):
        d, p = self._marr_setup('P000003', 'P000002')
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertEqual(s['citas_ligadas_a_evento'], 1)
        self.assertNotIn('EXTRACTO_SECRETO', o)

    def test_marr_citation_extract_kept_when_both_public(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        for r in (2, 3, 4):
            ws.cell(r, 6).value = '1850-01-01'; ws.cell(r, 8).value = '1920-01-01'
        wb['06_Documentos'].cell(5, 13).value = 'Publico'
        c = wb['05_Citas']; c.cell(5, 3).value = 'P000002'; c.cell(5, 4).value = 'Matrimonio'; c.cell(5, 5).value = '1975-01-01'; c.cell(5, 7).value = 'EXTRACTO_VISIBLE'
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), bind_events=True)
        o = self._out(d)
        self.assertIn('EXTRACTO_VISIBLE', o)
        self.assertIn('1 MARR', o)
        self.assertEqual(s['personas_privadas_RESN'], 0)

    def test_default_does_not_bind_events(self):
        d, p = self._marr_setup('P000002', 'P000003')
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertEqual(s['citas_ligadas_a_evento'], 0)
        self.assertEqual(s['citas_ligadas_a_persona'], 1)
        self.assertIn('PAGE Folio 1 | C000001 | Matrimonio: 1975-01-01', o)
        self.assertEqual(o.count('1 MARR'), 1)
        self.assertNotIn('2 SOUR @S000001@', o.split('0 @F1@ FAM')[1])

    def _cli(self, *args):
        import subprocess
        return subprocess.run([sys.executable, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'mapper.py')] + list(args), capture_output=True, text=True)

    def test_cli_rejects_malformed_invocations_without_writing(self):
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        for args in ([], [p, out, 'extra'], ['--bind-events', p, out], [p, '--bind-events', out], [p, out, '--bind-events'], [p, p], [p, '-x']):
            r = self._cli(*args)
            self.assertEqual(r.returncode, 2, args)
            self.assertFalse(os.path.exists(out), args)

    def test_cli_valid_invocation_and_no_binding_exposed(self):
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        r = self._cli(p, out)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue(os.path.exists(out))
        self.assertIn('citas_ligadas_a_evento: 0', r.stdout)

    def test_late_error_keeps_existing_output_and_writes_nothing_new(self):
        from unittest import mock
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        with open(out, 'w') as fh:
            fh.write('ANTERIOR\n')
        out2 = os.path.join(d, 'nuevo.ged')
        # fallos en la fase de publicacion (ultima): creacion del temporal y reemplazo
        for target in ('mapper.tempfile.mkstemp', 'mapper.os.replace'):
            for o in (out, out2):
                with mock.patch(target, side_effect=OSError('fallo tardio')):
                    with self.assertRaises(OSError):
                        mapper.run(p, o)
        with open(out) as fh:
            self.assertEqual(fh.read(), 'ANTERIOR\n')
        self.assertFalse(os.path.exists(out2))
        self.assertEqual([f for f in os.listdir(d) if f.startswith('.mapper-')], [])

    def test_failure_at_publication_keeps_existing_output(self):
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        with open(out, 'w') as fh:
            fh.write('ANTERIOR\n')
        import unittest.mock as um
        # tempfile=None hace fallar la publicacion (ultima fase): hasta ahi no debe haberse tocado la salida
        with um.patch.object(mapper, 'tempfile', None):
            with self.assertRaises(AttributeError):
                mapper.run(p, out)
        with open(out) as fh:
            self.assertEqual(fh.read(), 'ANTERIOR\n')

    def test_double_failure_replace_and_unlink_warns_and_keeps_original_error(self):
        import io, stat
        from unittest import mock
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        with open(out, 'w') as fh:
            fh.write('ANTERIOR\n')
        err = io.StringIO()
        with mock.patch('mapper.os.replace', side_effect=OSError('fallo replace')), \
             mock.patch('mapper.os.unlink', side_effect=OSError('fallo unlink')), \
             mock.patch('sys.stderr', err):
            with self.assertRaises(OSError) as cm:
                mapper.run(p, out)
        self.assertIn('fallo replace', str(cm.exception))
        left = [f for f in os.listdir(d) if f.startswith('.mapper-')]
        self.assertEqual(len(left), 1)
        path = os.path.join(d, left[0])
        self.assertIn(path, err.getvalue())
        self.assertEqual(stat.S_IMODE(os.stat(path).st_mode), 0o600)
        self.assertNotIn('Ficticia', err.getvalue())
        with open(out) as fh:
            self.assertEqual(fh.read(), 'ANTERIOR\n')

    def test_cli_rejects_aliases_of_input_without_mutating(self):
        import hashlib
        d, p, wb = self._fixture()
        sha = lambda f: hashlib.sha256(open(f, 'rb').read()).hexdigest()
        h0 = sha(p)
        sym = os.path.join(d, 'sym.ged'); os.symlink(p, sym)
        hard = os.path.join(d, 'hard.ged'); os.link(p, hard)
        broken = os.path.join(d, 'broken.ged'); os.symlink(os.path.join(d, 'nada.ged'), broken)
        brk_in = os.path.join(d, 'brk_in.ged'); os.symlink(p, brk_in)
        dirlink = os.path.join(d, 'sub'); os.mkdir(dirlink)
        alias_dir = os.path.join(d, 'aliasdir'); os.symlink(os.path.dirname(p), alias_dir)
        cases = [[p, sym], [p, hard], [p, broken], [sym, p], [p, os.path.join(alias_dir, os.path.basename(p))], [p, dirlink],
                 [p, os.path.join(os.path.dirname(p), '.', os.path.basename(p))]]
        for args in cases:
            r = self._cli(*args)
            self.assertEqual(r.returncode, 2, (args, r.stdout))
        self.assertEqual(sha(p), h0)
        self.assertTrue(os.path.islink(sym)); self.assertTrue(os.path.islink(broken))
        self.assertEqual(os.stat(hard).st_ino, os.stat(p).st_ino)
        self.assertEqual(sha(hard), h0)
        self.assertFalse(os.path.exists(os.path.join(d, 'nada.ged')))
        self.assertEqual([f for f in os.listdir(d) if f.startswith('.mapper-')], [])

    def test_library_run_also_rejects_aliases(self):
        d, p, wb = self._fixture()
        hard = os.path.join(d, 'hard.ged'); os.link(p, hard)
        sym = os.path.join(d, 'sym.ged'); os.symlink(p, sym)
        for o in (hard, sym, p):
            with self.assertRaises(mapper.MapperError):
                mapper.run(p, o)

    def test_absent_output_and_normal_existing_output_are_valid(self):
        d, p, wb = self._fixture()
        r = self._cli(p, os.path.join(d, 'nueva', 'x.ged'))
        self.assertEqual(r.returncode, 0, r.stderr)
        r = self._cli(p, os.path.join(d, 'nueva', 'x.ged'))
        self.assertEqual(r.returncode, 0, r.stderr)

    def _add_child(self, padre, madre, pid='P000090'):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        row = [None] * 16
        row[0], row[1], row[2], row[4], row[9], row[10] = pid, 'Nuevo', 'Hijo', 'M', padre, madre
        ws.append(row)
        wb.save(p)
        return d, p

    def test_role_conflicts_fail_closed_without_output(self):
        # en el fixture P000002 es el padre y P000003 la madre de P000001
        casos = {
            'misma persona padre y madre': ('P000002', 'P000002'),
            'persona como padre y como madre en hijos distintos': ('P000001', 'P000002'),
            'roles de la pareja intercambiados': ('P000003', 'P000002'),
        }
        for nombre, (pa, ma) in casos.items():
            d, p = self._add_child(pa, ma)
            out = os.path.join(d, 'out.ged')
            with open(out, 'w') as fh:
                fh.write('ANTERIOR\n')
            with self.assertRaises(mapper.MapperError, msg=nombre) as cm:
                mapper.run(p, out)
            self.assertIn('conflicto de roles', str(cm.exception), nombre)
            with open(out) as fh:
                self.assertEqual(fh.read(), 'ANTERIOR\n', nombre)
            r = self._cli(p, os.path.join(d, 'cli.ged'))
            self.assertEqual(r.returncode, 1, nombre)
            self.assertFalse(os.path.exists(os.path.join(d, 'cli.ged')), nombre)

    def test_self_spouse_fails_closed(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        ws.cell(3, 12).value = 'P000002'      # id_conyuge = si mismo
        wb.save(p)
        out = os.path.join(d, 'out.ged')
        with self.assertRaises(mapper.MapperError) as cm:
            mapper.run(p, out)
        self.assertIn('id_conyuge es la propia persona', str(cm.exception))
        self.assertFalse(os.path.exists(out))

    def test_spouse_row_does_not_override_roles_from_children(self):
        d, p, wb = self._fixture()
        ws = wb['30_Arbol_Confirmado']
        ws.cell(3, 12).value = 'P000003'; ws.cell(4, 12).value = 'P000002'   # conyuges declarados en ambos sentidos
        ws.cell(3, 5).value = ''; ws.cell(4, 5).value = ''                   # sin sexo: no se inventa
        wb.save(p)
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertEqual(s['familias'], 1)
        self.assertIn('1 HUSB @P000002@', o)      # roles tomados de id_padre/id_madre del hijo
        self.assertIn('1 WIFE @P000003@', o)
        self.assertFalse(any('sexo no determinante' in w for w in warn))

    def test_role_conflict_message_names_affected_records(self):
        d, p = self._add_child('P000003', 'P000002')
        with self.assertRaises(mapper.MapperError) as cm:
            mapper.run(p, os.path.join(d, 'o.ged'))
        m = str(cm.exception)
        self.assertIn('P000090', m); self.assertIn('P000001', m)

    def test_second_child_with_same_roles_shares_family(self):
        d, p = self._add_child('P000002', 'P000003')
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        o = self._out(d)
        self.assertEqual(s['familias'], 1)
        self.assertEqual(o.count('1 CHIL'), 2)
        self.assertEqual(o.count('1 HUSB'), 1)

    def _media(self, privacy='Público — acta >100 años (CNIL)', excel_hash='auto', old=True, manifest_sha=None, rel='media/D000001.pdf', did='D000001', content=b'%PDF-1.4 sintetico'):
        import hashlib
        d, p, wb = self._fixture()
        mdir = os.path.join(d, 'lote'); os.makedirs(os.path.join(mdir, 'media'))
        real = hashlib.sha256(content).hexdigest()
        with open(os.path.join(mdir, 'media', 'D000001.pdf'), 'wb') as fh:
            fh.write(content)
        ws = wb['30_Arbol_Confirmado']
        if old:
            ws.cell(2, 6).value = '1850-01-01'; ws.cell(2, 8).value = '1920-01-01'
        d6 = wb['06_Documentos']
        d6.cell(5, 13).value = privacy
        d6.cell(5, 14).value = real if excel_hash == 'auto' else excel_hash
        wb.save(p)
        man = os.path.join(mdir, 'manifiesto.csv')
        with open(man, 'w') as fh:
            fh.write('ID_Documento,ruta_relativa,sha256\n%s,%s,%s\n' % (did, rel, manifest_sha or real))
        return d, p, man, real

    def test_media_happy_path_emits_obje_with_relative_path(self):
        d, p, man, real = self._media()
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=man)
        o = self._out(d)
        self.assertIn('0 @M000001@ OBJE', o)
        self.assertIn('1 FILE media/D000001.pdf', o)
        self.assertIn('2 FORM pdf', o)
        self.assertIn('1 OBJE @M000001@', o)
        self.assertIn(real, o)
        self.assertEqual(s['media_objetos_enlazados'], 1)
        self.assertEqual(s['media_no_migrados_por_motivo'], {})
        self.assertNotIn('doc_ficticio', o)

    def test_without_manifest_no_media_objects(self):
        d, p, man, real = self._media()
        s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'))
        self.assertNotIn('OBJE', self._out(d))
        self.assertEqual(s['media_manifiesto'], 'no usado')

    def test_media_not_migrated_for_private_dni_living_or_missing_excel_hash(self):
        casos = [({'privacy': 'Privado'}, 'privacidad_no_publica'),
                 ({'privacy': 'Público (contiene nº de DNI; solo máster privado)'}, 'privacidad_no_publica'),
                 ({'privacy': 'Privado — CNIL: no publicar en línea (nacimiento <100 años)'}, 'privacidad_no_publica'),
                 ({'privacy': ''}, 'privacidad_no_publica'),
                 ({'old': False}, 'persona_posiblemente_viva'),
                 ({'excel_hash': ''}, 'sin_hash_valido_en_excel')]
        for kw, motivo in casos:
            d, p, man, real = self._media(**kw)
            s, warn, rep, unsup = mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=man)
            self.assertNotIn('OBJE', self._out(d), kw)
            self.assertEqual(s['media_no_migrados_por_motivo'], {motivo: 1}, kw)

    def test_media_manifest_fails_closed(self):
        import hashlib
        bad = '0' * 64
        casos = {
            'hash del manifiesto distinto': dict(manifest_sha=bad),
            'hash del Excel distinto': dict(excel_hash=bad),
            'ID desconocido': dict(did='D999999'),
            'ruta absoluta': dict(rel='/etc/passwd'),
            'ruta con ..': dict(rel='media/../../x.pdf'),
            'ruta con barra invertida': dict(rel='media\\D000001.pdf'),
            'archivo ausente': dict(rel='media/NO.pdf'),
        }
        for nombre, kw in casos.items():
            d, p, man, real = self._media(**kw)
            out = os.path.join(d, 'o.ged')
            with open(out, 'w') as fh:
                fh.write('ANTERIOR\n')
            with self.assertRaises(mapper.MapperError, msg=nombre):
                mapper.run(p, out, media_manifest=man)
            with open(out) as fh:
                self.assertEqual(fh.read(), 'ANTERIOR\n', nombre)

    def test_media_manifest_symlink_duplicate_and_header_fail_closed(self):
        d, p, man, real = self._media()
        base = os.path.dirname(man)
        os.symlink(os.path.join(base, 'media', 'D000001.pdf'), os.path.join(base, 'media', 'link.pdf'))
        with open(man, 'w') as fh:
            fh.write('ID_Documento,ruta_relativa,sha256\nD000001,media/link.pdf,%s\n' % real)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=man)
        with open(man, 'w') as fh:
            fh.write('ID_Documento,ruta_relativa,sha256\nD000001,media/D000001.pdf,%s\nD000001,media/D000001.pdf,%s\n' % (real, real))
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=man)
        with open(man, 'w') as fh:
            fh.write('id,ruta,hash\nD000001,media/D000001.pdf,%s\n' % real)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=man)
        with self.assertRaises(mapper.MapperError):
            mapper.run(p, os.path.join(d, 'o.ged'), media_manifest=os.path.join(base, 'no_existe.csv'))
        self.assertFalse(os.path.exists(os.path.join(d, 'o.ged')))

    def test_cli_media_manifest_option_strict(self):
        d, p, man, real = self._media()
        out = os.path.join(d, 'o.ged')
        r = self._cli(p, out, '--media-manifest=' + man)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn('1 FILE media/D000001.pdf', self._out(d))
        os.unlink(out)
        for args in ([p, out, '--media-manifest=' + man, '--media-manifest=' + man], [p, out, '--media-manifest='], [p, out, '--media-manifest', man], [p, out, '--otra=1']):
            r = self._cli(*args)
            self.assertEqual(r.returncode, 2, args)
            self.assertFalse(os.path.exists(out), args)
        r = self._cli(p, out, '--media-manifest=' + os.path.join(d, 'nada.csv'))
        self.assertEqual(r.returncode, 1)
        self.assertFalse(os.path.exists(out))

    def test_successful_run_replaces_output_without_temp_leftovers(self):
        d, p, wb = self._fixture()
        out = os.path.join(d, 'out.ged')
        open(out, 'w').write('ANTERIOR\n')
        mapper.run(p, out)
        self.assertTrue(open(out).read().startswith('0 HEAD'))
        self.assertEqual([f for f in os.listdir(d) if f.startswith('.mapper-')], [])


if __name__ == '__main__':
    unittest.main()
