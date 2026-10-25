"""Excel genealogico -> GEDCOM 5.5.1 (solo lectura sobre el Excel, sin red)."""
import sys, os, re, datetime, collections, tempfile, csv, hashlib
import openpyxl

MM = ['JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP','OCT','NOV','DEC']
CONF = {'Muy alta': 3, 'Alta': 2, 'Baja': 1}   # QUAY; Gramps importa 0,1,2,3 como 0,1,3,4; sin QUAY = Normal(2)
NORMAL_PLAIN = ('Media',)
EVENT_MAP = {'Nacimiento': 'BIRT', 'Acta de nacimiento': 'BIRT', 'Defunción': 'DEAT', 'Acta de defunción': 'DEAT', 'Matrimonio': 'MARR', 'Acta de matrimonio': 'MARR'}   # valores EXACTOS de Tipo_hecho; el resto queda ligado a la persona
NORMAL_CONVERTED = ('Media-alta',)              # se importa como Normal con nota de conversion
MAXLEN = 240


class MapperError(Exception):
    pass


def nm(x):
    return re.sub(r'\s+', ' ', str(x if x is not None else '')).strip()


def cv(x):
    """Valor de celda para notas: vacio -> 'n/d' (nunca el literal None)."""
    t = nm(x)
    return t if t else 'n/d'


XREF = re.compile(r'^[A-Za-z0-9_]{1,20}$')
REQUIRED = {
    '30_Arbol_Confirmado': (0, ['id_persona', 'nombre', 'apellido1', 'apellido2', 'sexo', 'fecha_nacimiento', 'lugar_nacimiento', 'fecha_defuncion', 'lugar_defuncion', 'id_padre', 'id_madre', 'id_conyuge', 'fecha_matrimonio', 'lugar_matrimonio', 'fuente', 'notas']),
    '02_Relaciones': (3, ['ID_Persona_1', 'Tipo_relacion', 'ID_Persona_2']),
    '04_Fuentes': (3, ['ID_Fuente', 'Titulo', 'Autor_o_entidad', 'Tipo_fuente', 'Original_o_derivada', 'Calidad_fuente', 'Repositorio', 'Signatura', 'URL']),
    '05_Citas': (3, ['ID_Cita', 'ID_Fuente', 'ID_Persona', 'Tipo_hecho', 'Valor_afirmado', 'Localizador_pagina_folio', 'Extracto_o_resumen', 'Evaluacion_evidencia', 'Corroborada', 'Confianza', 'ID_Documento', 'URL']),
    '06_Documentos': (3, ['ID_Documento', 'Nombre_archivo', 'Tipo_documento', 'ID_Persona_principal', 'ID_Fuente', 'ID_Cita', 'Derechos_privacidad', 'SHA256']),
}
PUBLIC_OK = ('publico', 'público', 'public')
CURRENT_YEAR_LIMIT = 1906   # nacidos despues y sin defuncion registrada: se tratan como posiblemente vivos


def check_xref(kind, ident):
    if not XREF.match(str(ident)):
        raise MapperError(f'{kind} con formato de XREF invalido: {ident!r}')


def gedcom_text(level, tag, text):
    """Linea GEDCOM con CONC/CONT segun la especificacion: sin truncar nada."""
    out = []
    parts = str(text).replace('\r\n', '\n').replace('\r', '\n').split('\n')
    first = True
    for part in parts:
        chunks = [part[i:i + MAXLEN] for i in range(0, len(part), MAXLEN)] or ['']
        for j, ch in enumerate(chunks):
            ult = j == len(chunks) - 1                  # solo el ultimo trozo pierde espacios finales: en trozos intermedios un espacio es contenido (CONC lo une sin separador)
            if first:
                out.append((f'{level} {tag} {ch}'.rstrip() if ult else f'{level} {tag} {ch}') if ch else f'{level} {tag}')
                first = False
            elif j == 0:
                out.append(f'{level + 1} CONT {ch}'.rstrip() if ult else f'{level + 1} CONT {ch}')
            else:
                out.append(f'{level + 1} CONC {ch}')
    return out


def parse_date(s, warn=None, ctx=''):
    """Devuelve (fecha_gedcom|None, nota_original|None). Si no se puede mapear:
    (None, texto original) para que la nota original siempre se conserve."""
    s = nm(s)
    if not s:
        return None, None

    def valid(y, m=1, d=1):
        try:
            datetime.date(int(y), int(m), int(d))
            return True
        except ValueError:
            return False
    m = re.match(r'c\. (\d{4})-(\d{4})$', s)
    if m and valid(m[1]) and valid(m[2]) and int(m[1]) <= int(m[2]):
        return f'BET {m[1]} AND {m[2]}', s
    m = re.match(r'c\. (\d{4})$', s)
    if m and valid(m[1]):
        return f'ABT {m[1]}', s
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m and valid(*m.groups()):
        return f'{int(m[3])} {MM[int(m[2]) - 1]} {m[1]}', None
    m = re.match(r'(\d{4})-(\d{2})$', s)
    if m and valid(m[1], m[2]):
        return f'{MM[int(m[2]) - 1]} {m[1]}', None
    if re.match(r'\d{4}$', s) and valid(s):
        return s, None
    if warn is not None:
        warn.append(f'fecha no mapeada ({ctx}): {s}')
    return None, s


def table(wb, name, header_row=3):
    rows = list(wb[name].iter_rows(values_only=True))
    hdr = rows[header_row]
    return [dict(zip(hdr, r)) for r in rows[header_row + 1:] if r[0]]


def unique(items, key, what):
    seen = set()
    for it in items:
        k = it[key]
        if k in seen:
            raise MapperError(f'{what} duplicado: {k}')
        seen.add(k)


LOC_UNSAFE = re.compile(r'\b(DNI|NIE|pasaporte|passport)\b|\d{8}|@|https?:', re.I)


def loc_safe(v):
    """Localizador de cita de persona privada: se conserva solo si es una referencia documental corta sin DNI/NIE/pasaporte, 8+ digitos seguidos, @ ni URL."""
    s = nm(v)
    return bool(s) and len(s) <= 120 and not LOC_UNSAFE.search(s)


# --- Texto libre (notas, evaluacion, extracto, titulos...): un unico control para todos los campos -------------
FREE_IDENTITY = re.compile(r'\b(DNI|NIE|pasaporte|passport|carnet|licencia de conducir|permiso de conducir|libro de familia|documento nacional de identidad)\b', re.I)
URL_ANY = re.compile(r'(?i)\b(?:https?|ftp|file)://[^\s<>"\')\]]+|\bwww\.[^\s<>"\')\]]+')
EMAIL_ANY = re.compile(r'[\w.+-]+@[\w-]+(?:\.[\w-]+)+')
LONG_DIGITS = re.compile(r'(?<![\w])\d{8,}(?![\w])')
# Hosts de almacenamiento/mensajeria/enlaces compartibles: nunca se emiten. Los archivos y sitios genealogicos publicos (archivos, Geneanet, etc.) se conservan.
PRIVATE_HOSTS = ('drive.google.com', 'docs.google.com', 'mail.google.com', 'sheets.google.com', 'photos.google.com', 'photos.app.goo.gl', 'goo.gl', 'forms.gle',
                 'dropbox.com', 'dropboxusercontent.com', '1drv.ms', 'onedrive.live.com', 'sharepoint.com', 'icloud.com', 'box.com', 'wetransfer.com', 'we.tl', 'mega.nz', 'mega.io',
                 'notion.so', 'notion.site', 'whatsapp.com', 'wa.me', 't.me', 'telegram.me', 'bit.ly', 'tinyurl.com', 'localhost')
SECRET_QUERY = re.compile(r'(?i)(^|[?&;])(token|access_token|key|api_?key|sig|signature|auth|authuser|share|sharing|usp|invite|session|sid|password|pwd)=')


def url_private(u):
    """True si la URL es un enlace privado/compartible, sin HTTPS, local, IP, o lleva credenciales/token en la consulta."""
    from urllib.parse import urlsplit
    raw = u if re.match(r'(?i)[a-z][a-z0-9+.-]*://', u) else 'https://' + u
    try:
        sp = urlsplit(raw)
        host = (sp.hostname or '').lower()
    except ValueError:
        return True
    if sp.scheme.lower() != 'https' or not host or sp.username or sp.password:
        return True
    if re.fullmatch(r'[\d.]+|[0-9a-f:]+', host) or '.' not in host:
        return True
    if any(host == h or host.endswith('.' + h) for h in PRIVATE_HOSTS):
        return True
    return bool(SECRET_QUERY.search(sp.query) or SECRET_QUERY.search(sp.fragment))


def scrub(text, ctx, omis):
    """Control comun de texto libre emitido en campos NO privados. Devuelve el texto saneado y registra contadores por ctx:
    referencia a documento de identidad -> campo entero omitido; enlace privado -> solo ese enlace; correo y numeros de 8+ cifras -> solo ese fragmento."""
    s = nm(text)
    if not s:
        return s
    if FREE_IDENTITY.search(s):
        omis['libre_identidad'].append(ctx)
        return '[campo omitido por privacidad: referencia a documento de identidad (ver Excel)]'
    def url(m):
        if url_private(m.group(0)):
            omis['libre_enlace_privado'].append(ctx)
            return '[enlace omitido por privacidad]'
        return m.group(0)
    s = URL_ANY.sub(url, s)
    if EMAIL_ANY.search(s):
        omis['libre_correo'].append(ctx)
        s = EMAIL_ANY.sub('[correo omitido por privacidad]', s)
    if LONG_DIGITS.search(s):
        omis['libre_numero_largo'].append(ctx)
        s = LONG_DIGITS.sub('[numero omitido por privacidad]', s)
    return s


SRC_IDENTITY = re.compile(r'\b(DNI|NIE|pasaporte|passport|carnet|licencia de conducir|permiso de conducir|libro de familia)\b', re.I)
URL_UNSAFE = re.compile(r'\b(DNI|NIE|pasaporte|passport)\b|\d{8}|@', re.I)
# Sources whose title names specific people: restricted by default (example ids).
SRC_RESTRICT_DEFAULT = {'S000112', 'S000209', 'S000210', 'S000093', 'S000039'}   # default restricted source ids (example values)
LOC_OMIT_DEFAULT = {'C000188', 'C000091', 'C000094', 'C000097', 'C000117'}   # default omitted locator citation ids (example values)
SRC_PERSONAL = re.compile(r'testimonio|declaraci|whatsapp|conversaci|correo|gmail|mail\.google|e-?mail|mensaje', re.I)
MEDIA_FORBIDDEN = ('privado', 'no publicar', 'dni', '<100', 'expediente')


def media_public(v):
    """Un documento es migrable como archivo solo si su privacidad dice Publico/Publica y no contiene marcas restrictivas."""
    s = nm(v).lower()
    return (s.startswith('públic') or s.startswith('public')) and not any(w in s for w in MEDIA_FORBIDDEN)


def load_media_manifest(manifest_path, doc_ids, excel_hash):
    """CSV ID_Documento,ruta_relativa,sha256 -> {ID: (ruta_relativa, sha256)}. Fail closed: ID desconocido/duplicado, ruta absoluta,
    con '..', '\\' o symlink, archivo ausente, hash del manifest distinto del archivo, o hash del Excel (si es valido) distinto."""
    if not os.path.isfile(manifest_path):
        raise MapperError(f'no existe el manifiesto de medios: {manifest_path}')
    base = os.path.dirname(os.path.realpath(manifest_path))
    out = {}
    with open(manifest_path, newline='', encoding='utf-8') as fh:
        rd = csv.DictReader(fh)
        if not rd.fieldnames or [h.strip() for h in rd.fieldnames] != ['ID_Documento', 'ruta_relativa', 'sha256']:
            raise MapperError('manifiesto de medios: la cabecera debe ser exactamente ID_Documento,ruta_relativa,sha256')
        for n, r in enumerate(rd, 2):
            did, rel, sha = (r['ID_Documento'] or '').strip(), (r['ruta_relativa'] or '').strip(), (r['sha256'] or '').strip().lower()
            if did not in doc_ids:
                raise MapperError(f'manifiesto de medios linea {n}: ID_Documento desconocido {did!r}')
            if did in out:
                raise MapperError(f'manifiesto de medios linea {n}: ID_Documento duplicado {did}')
            if not rel or rel.startswith('/') or '\\' in rel or re.search(r'(^|/)\.\.(/|$)', rel) or re.match(r'^[A-Za-z]:', rel):
                raise MapperError(f'manifiesto de medios linea {n}: ruta no relativa o insegura {rel!r}')
            if not re.fullmatch(r'[0-9a-f]{64}', sha):
                raise MapperError(f'manifiesto de medios linea {n}: sha256 no valido para {did}')
            full = os.path.join(base, rel)
            if os.path.islink(full) or not os.path.isfile(full) or os.path.realpath(full) != os.path.join(base, os.path.normpath(rel)):
                raise MapperError(f'manifiesto de medios linea {n}: archivo ausente, enlace o fuera de la carpeta: {rel}')
            h = hashlib.sha256()
            with open(full, 'rb') as fh2:
                for blk in iter(lambda: fh2.read(1 << 20), b''):
                    h.update(blk)
            if h.hexdigest() != sha:
                raise MapperError(f'manifiesto de medios linea {n}: el sha256 del archivo no coincide con el manifiesto ({did})')
            ex = excel_hash.get(did, '')
            if re.fullmatch(r'[0-9a-fA-F]{64}', ex) and ex.lower() != sha:
                raise MapperError(f'{did}: el sha256 del Excel no coincide con el archivo')
            out[did] = (rel, sha)
    return out


def check_output_path(inp, out):
    """Rechaza una salida que sea la entrada por cualquier alias (ruta, symlink, hardlink) o un symlink (os.replace cambiaria el enlace, no su destino).
    Salida ausente: valido. Symlink roto: rechazado. Devuelve un mensaje de error o None."""
    if os.path.islink(out):
        return f'la salida es un enlace simbolico ({out}); indica una ruta de archivo normal'
    if os.path.realpath(inp) == os.path.realpath(out):
        return 'la salida no puede ser el archivo de entrada'
    try:
        if os.path.exists(out) and os.path.samefile(inp, out):
            return 'la salida no puede ser el archivo de entrada (mismo archivo, enlace duro)'
    except OSError as ex:
        return f'no se pudo comprobar la salida: {ex}'
    if os.path.isdir(out):
        return f'la salida es un directorio: {out}'
    return None


def run(path, out_path, bind_events=False, media_manifest=None, persons=None, rels=None, evidencia=None):
    if not os.path.isfile(path):
        raise MapperError(f'no existe el archivo: {path}')
    _msg = check_output_path(path, out_path)
    if _msg:
        raise MapperError(_msg)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sh, (hr, cols) in REQUIRED.items():
        if persons is not None and sh == '30_Arbol_Confirmado':
            continue  # V2: las personas vienen del selector (01+02), no de la hoja 30
        if sh not in wb.sheetnames:
            raise MapperError(f'falta la hoja {sh}')
        hdr = list(wb[sh].iter_rows(min_row=hr + 1, max_row=hr + 1, values_only=True))[0]
        miss = [c for c in cols if c not in hdr]
        if miss:
            raise MapperError(f'hoja {sh}: faltan columnas {miss}')
    warn, stats = [], collections.Counter()
    if persons is not None:
        P = persons
    else:
        r30 = list(wb['30_Arbol_Confirmado'].iter_rows(values_only=True))
        P = [dict(zip(r30[0], r)) for r in r30[1:] if r[0]]
    unique(P, 'id_persona', 'id_persona')
    byid = {p['id_persona']: p for p in P}
    R = rels if rels is not None else table(wb, '02_Relaciones')
    Sl = table(wb, '04_Fuentes'); unique(Sl, 'ID_Fuente', 'ID_Fuente')
    S = {s['ID_Fuente']: s for s in Sl}
    C = table(wb, '05_Citas'); unique(C, 'ID_Cita', 'ID_Cita')
    D = table(wb, '06_Documentos'); unique(D, 'ID_Documento', 'ID_Documento')
    media = load_media_manifest(media_manifest, {x['ID_Documento'] for x in D}, {x['ID_Documento']: nm(x['SHA256']) for x in D}) if media_manifest else {}
    media_used = {}
    media_reasons = collections.Counter()
    dup01 = {}
    if '01_Personas' in wb.sheetnames:
        cnt = collections.Counter(r[0] for r in wb['01_Personas'].iter_rows(min_row=5, values_only=True) if r[0])
        dup01 = {k: v for k, v in cnt.items() if v > 1}
    cnt02 = collections.Counter(r['ID_Relacion'] for r in R if r.get('ID_Relacion'))
    dup02 = {k: v for k, v in cnt02.items() if v > 1}
    if dup02:
        warn.append(f'02_Relaciones con ID_Relacion duplicado (no afecta a la conversion, a revisar): {sorted(dup02)}')
    refs = {p['id_persona'] for p in P} | {p[k] for p in P for k in ('id_padre', 'id_madre', 'id_conyuge') if p[k]}
    hit = sorted(refs & set(dup01))
    if hit:
        raise MapperError(f'IDs usados por el arbol confirmado estan duplicados en 01_Personas (personas distintas?): {hit}')
    if dup01:
        warn.append(f'01_Personas con IDs duplicados fuera del arbol confirmado: {sorted(dup01)}')
    for p in P: check_xref('id_persona', p['id_persona'])
    for x in Sl: check_xref('ID_Fuente', x['ID_Fuente'])
    clash = set(byid) & set(S)
    if clash:
        raise MapperError(f'colision de XREF entre personas y fuentes: {sorted(clash)[:5]}')
    for ref in ('id_padre', 'id_madre', 'id_conyuge'):
        for p in P:
            if p[ref] and p[ref] not in byid:
                raise MapperError(f'{p["id_persona"]}: {ref}={p[ref]!r} no esta en la hoja 30')

    # Familias: clave = conjunto no ordenado de progenitores (sin depender del sexo)
    fams = {}

    def fam(key):
        return fams.setdefault(key, {'husb': None, 'wife': None, 'ch': [], 'mar': None, 'mars': []})
    # Conflictos de rol: se rechazan (fail closed) en vez de inventar o pisar HUSB/WIFE
    como_padre, como_madre, roles_pareja, role_conf = {}, {}, {}, []
    for p in P:
        f, m = p['id_padre'] or '', p['id_madre'] or ''
        if f and f == m:
            role_conf.append(f"{p['id_persona']}: id_padre e id_madre son la misma persona ({f})")
        if f:
            como_padre.setdefault(f, []).append(p['id_persona'])
        if m:
            como_madre.setdefault(m, []).append(p['id_persona'])
        if f and m and f != m:
            roles_pareja.setdefault(frozenset((f, m)), {}).setdefault((f, m), []).append(p['id_persona'])
    for x in sorted(set(como_padre) & set(como_madre)):
        role_conf.append(f"{x}: figura como padre de {sorted(como_padre[x])} y como madre de {sorted(como_madre[x])}")
    for k, v in roles_pareja.items():
        if len(v) > 1:
            role_conf.append('pareja ' + '/'.join(sorted(k)) + ': roles padre/madre intercambiados entre hijos ' + '; '.join(f"{a}=padre,{b}=madre en {sorted(h)}" for (a, b), h in sorted(v.items())))
    for p in P:
        if p['id_conyuge'] and p['id_conyuge'] == p['id_persona']:
            role_conf.append(f"{p['id_persona']}: id_conyuge es la propia persona")
    if role_conf:
        raise MapperError('conflicto de roles padre/madre en hoja 30 (no se asignan roles; corregir en el Excel): ' + ' | '.join(role_conf))
    for p in P:
        f, m = p['id_padre'] or '', p['id_madre'] or ''
        if f or m:
            e = fam(frozenset(x for x in (f, m) if x))
            e['husb'], e['wife'] = (f or e['husb']), (m or e['wife'])
            e['ch'].append(p['id_persona'])
    for p, c, c_fecha, c_lugar in [(p, c, f, l) for p in P for (c, f, l) in (p['conyuges'] if 'conyuges' in p else [(p['id_conyuge'], p['fecha_matrimonio'], p['lugar_matrimonio'])])]:
        if c in byid:
            key = frozenset((p['id_persona'], c)); e = fam(key)
            if not (e['husb'] and e['wife']):
                a, b = p['id_persona'], c
                sx = {a: byid[a]['sexo'], b: byid[b]['sexo']}
                if sx[a] == 'M' and sx[b] != 'M':
                    e['husb'], e['wife'] = a, b
                elif sx[b] == 'M' and sx[a] != 'M':
                    e['husb'], e['wife'] = b, a
                elif sx[a] == 'F' and sx[b] != 'F':
                    e['husb'], e['wife'] = b, a
                elif sx[b] == 'F' and sx[a] != 'F':
                    e['husb'], e['wife'] = a, b
                else:
                    e['husb'], e['wife'] = sorted((a, b))
                    warn.append(f'sexo no determinante en pareja {a}/{b}: roles HUSB/WIFE por orden de ID, no por sexo')
            if c_fecha or c_lugar:
                e['mars'].append((nm(c_fecha), nm(c_lugar)))
    for key, e in fams.items():
        uniq = sorted(set(e['mars']))
        if len(uniq) > 1:
            warn.append(f'{sorted(key)}: datos de matrimonio contradictorios entre conyuges {uniq}; no se escribe MARR, quedan en nota')
            e['mar'] = None; e['mar_conflict'] = uniq
        elif uniq:
            e['mar'] = uniq[0]
    fid = {k: f'@F{i + 1}@' for i, k in enumerate(fams)}

    cit = collections.defaultdict(list)
    for c in C:
        if c['ID_Persona'] in byid:
            cit[c['ID_Persona']].append(c)

    out = ['0 HEAD', '1 SOUR PILOTO', '1 GEDC', '2 VERS 5.5.1', '1 CHAR UTF-8']
    used = set()
    fam_flags = {}
    for key, e in fams.items():
        no_union_ = bool(e['husb'] and e['wife']) and not any(r['Tipo_relacion'] in ('Cónyuge', 'Esposo', 'Esposa', 'Pareja (unión libre)') and {r['ID_Persona_1'], r['ID_Persona_2']} == {e['husb'], e['wife']} for r in R)
        fam_flags[key] = {'no_union': no_union_, 'writes_marr': bool(e['mar']) and not no_union_}
    fam_bound = collections.defaultdict(list)

    def privacy_of(p):
        docs = [x for x in D if x['ID_Persona_principal'] == p['id_persona']]
        by_ = re.match(r'(?:c\. )?(\d{4})', nm(p['fecha_nacimiento']))
        no_death_ = not (nm(p['fecha_defuncion']) or nm(p['lugar_defuncion']))
        living = no_death_ and (not by_ or int(by_[1]) > CURRENT_YEAR_LIMIT)
        dpriv = any(nm(x['Derechos_privacidad']).lower() not in PUBLIC_OK for x in docs)
        return living, dpriv
    priv_map = {q['id_persona']: any(privacy_of(q)) for q in P}
    reasons = collections.Counter()

    def cite_valid(c):
        return c['ID_Fuente'] in S and (c['Confianza'] in CONF or c['Confianza'] in NORMAL_PLAIN or c['Confianza'] in NORMAL_CONVERTED)

    omis = collections.defaultdict(list)
    sid_priv = set()

    def cite_lines(c, level, private):
        sid, cf = c['ID_Fuente'], c['Confianza']
        used.add(sid)
        if private:
            sid_priv.add(sid)
        o = [f'{level} SOUR @{sid}@']
        if private:
            loc = cv(c['Localizador_pagina_folio'])
            if not loc_safe(c['Localizador_pagina_folio']) or c['ID_Cita'] in LOC_OMIT_DEFAULT:
                if nm(c['Localizador_pagina_folio']):
                    omis['localizador'].append(c['ID_Cita'])
                loc = 'localizador omitido por privacidad'
            o += gedcom_text(level + 1, 'PAGE', nm(f"{loc} | {c['ID_Cita']} | {scrub(cv(c['Tipo_hecho']), c['ID_Cita'] + ':Tipo_hecho', omis)}"))
        else:
            cx = c['ID_Cita']
            o += gedcom_text(level + 1, 'PAGE', nm(f"{scrub(cv(c['Localizador_pagina_folio']), cx + ':Localizador', omis)} | {cx} | {scrub(cv(c['Tipo_hecho']), cx + ':Tipo_hecho', omis)}: {scrub(cv(c['Valor_afirmado']), cx + ':Valor_afirmado', omis)}"))
        if cf in CONF:
            o.append(f'{level + 1} QUAY {CONF[cf]}')
        conv = ' (convertida a Normal: sin equivalente exacto)' if cf in NORMAL_CONVERTED else ''
        if private:
            omis['cita_campos'].append(c['ID_Cita'])
            o += gedcom_text(level + 1, 'NOTE', nm(f"Confianza Excel: {cv(cf)}{conv}; Corroborada: {cv(c['Corroborada'])}; Valor afirmado, Evaluacion, Doc, URL y Extracto omitidos por privacidad (ver Excel)"))
        else:
            o += gedcom_text(level + 1, 'NOTE', nm(f"Confianza Excel: {cv(cf)}{conv}; Corroborada: {cv(c['Corroborada'])}; Evaluacion: {scrub(cv(c['Evaluacion_evidencia']), cx + ':Evaluacion', omis)}; Doc: {cv(c['ID_Documento'])}; URL: {scrub(cv(c['URL']), cx + ':URL', omis)}; Extracto: {scrub(cv(c['Extracto_o_resumen']), cx + ':Extracto', omis)}"))
        return o

    def years_of(text):
        return {int(y) for y in re.findall(r'(?<!\d)(\d{4})(?!\d)', str(text or ''))}

    def compatible(c, ev_years):
        vy = years_of(c['Valor_afirmado'])
        if not vy:
            return True, ''
        if not ev_years:
            return False, 'evento sin año para contrastar con el valor de la cita'
        if vy & ev_years or any(min(ev_years) <= y <= max(ev_years) for y in vy):
            return True, ''
        return False, 'año del valor de la cita no concuerda con el evento'
    for p in P:
        i = p['id_persona']
        out.append(f'0 @{i}@ INDI')
        surn = nm(p['apellido1']) + (' ' + nm(p['apellido2']) if p['apellido2'] else '')
        out.append(f"1 NAME {nm(p['nombre'])} /{surn}/")
        if p['sexo'] in ('M', 'F'):
            out.append('1 SEX ' + p['sexo'])
        else:
            warn.append(f'{i}: sexo no M/F ({p["sexo"]!r}); no se escribe SEX')
        docs_p = [x for x in D if x['ID_Persona_principal'] == i]
        by = re.match(r'(?:c\. )?(\d{4})', nm(p['fecha_nacimiento']))
        no_death = not (nm(p['fecha_defuncion']) or nm(p['lugar_defuncion']))
        maybe_living = no_death and (not by or int(by[1]) > CURRENT_YEAR_LIMIT)
        doc_private = any(nm(x['Derechos_privacidad']).lower() not in PUBLIC_OK for x in docs_p)
        private = maybe_living or doc_private
        bound = {'BIRT': [], 'DEAT': []}
        bound_ids = set()
        ev_info = {}
        for tag, df, pf in (('BIRT', 'fecha_nacimiento', 'lugar_nacimiento'), ('DEAT', 'fecha_defuncion', 'lugar_defuncion')):
            g0, n0 = parse_date(p[df])
            ev_info[tag] = (bool(g0 or n0 or p[pf]), years_of(g0) | years_of(n0))
        for c in cit[i]:
            if not cite_valid(c):
                continue
            ev = EVENT_MAP.get(nm(c['Tipo_hecho']))
            if not bind_events:
                reasons['ligado_a_persona_por_defecto'] += 1
                continue
            if not ev:
                reasons['tipo_no_exacto'] += 1
                continue
            if ev in ('BIRT', 'DEAT'):
                present, yrs = ev_info[ev]
                if not present:
                    reasons['evento_ausente'] += 1; continue
                ok, why = compatible(c, yrs)
                if not ok:
                    reasons[why] += 1; continue
                bound[ev].append(c); bound_ids.add(c['ID_Cita'])
            else:
                cands = [k for k, e in fams.items() if i in (e['husb'], e['wife']) and fam_flags[k]['writes_marr']]
                if len(cands) != 1:
                    reasons['matrimonio_sin_familia_unica_confirmada'] += 1; continue
                e = fams[cands[0]]
                ok, why = compatible(c, years_of(parse_date(e['mar'][0])[0]) | years_of(parse_date(e['mar'][0])[1]))
                if not ok:
                    reasons['matrimonio_' + why] += 1; continue
                spouses = [e['husb'], e['wife']]
                fam_private = private or any(priv_map.get(sp, True) for sp in spouses if sp) or not all(spouses)
                fam_bound[cands[0]].append((c, fam_private)); bound_ids.add(c['ID_Cita'])
        for tag, df, pf in (('BIRT', 'fecha_nacimiento', 'lugar_nacimiento'), ('DEAT', 'fecha_defuncion', 'lugar_defuncion')):
            g, note = parse_date(p[df], warn, f'{i} {tag}')
            if g or note or p[pf]:
                out.append('1 ' + tag)
                if g:
                    out.append('2 DATE ' + g)
                if p[pf]:
                    out.append('2 PLAC ' + nm(p[pf]))
                if note:
                    out += gedcom_text(2, 'NOTE', 'Fecha original en el Excel: ' + note + (' (no mapeada a fecha GEDCOM)' if not g else ' (rango/aproximada, no fecha exacta)'))
                for c in bound[tag]:
                    out += cite_lines(c, 2, private)
                    stats['citas_mapeadas'] += 1; stats['citas_a_evento'] += 1
        for key, e in fams.items():
            if i in key:
                out.append(f'1 FAMS {fid[key]}')
            if i in e['ch']:
                out.append(f'1 FAMC {fid[key]}')
        if private:
            out.append('1 RESN privacy')
            stats['personas_privadas'] += 1
            warn.append(f"{i}: marcada RESN privacy ({'posiblemente viva' if maybe_living else ''}{' y ' if maybe_living and doc_private else ''}{'documento no publico/desconocido' if doc_private else ''}); sin extractos ni nombres de archivo en notas")
        if p['notas']:
            if private:
                omis['nota_personal'].append(i)
                out += gedcom_text(1, 'NOTE', 'Nota personal omitida por privacidad (ver Excel)')
            else:
                out += gedcom_text(1, 'NOTE', scrub(p['notas'], i + ':Notas', omis))
        for x in docs_p:
            sha = nm(x['SHA256'])
            if re.fullmatch(r'[0-9a-fA-F]{64}', sha):
                sha_txt = 'SHA256 ' + sha
            else:
                sha_txt = 'SHA256 no acreditado' + (f" (valor no valido: {scrub(sha, x['ID_Documento'] + ':SHA256', omis)})" if sha else ' (vacio)')
                warn.append(f"{x['ID_Documento']}: hash no valido o vacio")
                stats['docs_sin_hash_valido'] += 1
            fn = 'archivo omitido por privacidad' if (private or nm(x['Derechos_privacidad']).lower() not in PUBLIC_OK) else scrub(cv(x['Nombre_archivo']), x['ID_Documento'] + ':Nombre_archivo', omis)
            did = x['ID_Documento']
            if media_manifest:
                if did not in media:
                    media_reasons['sin_archivo_en_manifiesto'] += 1
                elif not media_public(x['Derechos_privacidad']):
                    media_reasons['privacidad_no_publica'] += 1
                elif maybe_living:
                    media_reasons['persona_posiblemente_viva'] += 1
                elif not re.fullmatch(r'[0-9a-fA-F]{64}', sha):
                    media_reasons['sin_hash_valido_en_excel'] += 1
                else:
                    mid = 'M' + re.sub(r'\D', '', did)
                    check_xref('media', mid)
                    if mid in byid or mid in S:
                        raise MapperError(f'colision de XREF de medios con persona/fuente: {mid}')
                    if mid in media_used and media_used[mid] != did:
                        raise MapperError(f'colision de XREF de medios: {mid}')
                    media_used[mid] = did
                    out.append(f'1 OBJE @{mid}@')
                    stats['media_objetos_enlazados'] += 1
            if private or nm(x['Derechos_privacidad']).lower() not in PUBLIC_OK:
                omis['documento_campos'].append(x['ID_Documento'])
                tipo_t, priv_t = 'omitido por privacidad (ver Excel)', 'omitida por privacidad (ver Excel)'
            else:
                tipo_t, priv_t = scrub(cv(x['Tipo_documento']), x['ID_Documento'] + ':Tipo_documento', omis), cv(x['Derechos_privacidad'])
            out += gedcom_text(1, 'NOTE', f"Documento {x['ID_Documento']}: {fn} | tipo {tipo_t} | fuente {cv(x['ID_Fuente'])} | cita {cv(x['ID_Cita'])} | privacidad {priv_t} | {sha_txt}")
        for c in cit[i]:
            sid, cf = c['ID_Fuente'], c['Confianza']
            if sid not in S:
                out += gedcom_text(1, 'NOTE', f"Cita {c['ID_Cita']} EXCLUIDA: fuente {cv(sid)} no esta en 04_Fuentes")
                warn.append(f"{c['ID_Cita']}: fuente {sid} ausente en 04_Fuentes"); stats['excluidas_sin_fuente'] += 1
                continue
            if cf not in CONF and cf not in NORMAL_PLAIN and cf not in NORMAL_CONVERTED:
                out += gedcom_text(1, 'NOTE', f"Cita {c['ID_Cita']} EXCLUIDA: confianza original desconocida/invalida ({cv(cf)}); pendiente de verificar")
                warn.append(f"{c['ID_Cita']}: excluida (confianza {cf!r})"); stats['excluidas_confianza'] += 1
                continue
            if c['ID_Cita'] in bound_ids:
                continue
            out += cite_lines(c, 1, private)
            stats['citas_a_persona'] += 1
            stats['citas_mapeadas'] += 1
    unsup = []
    ev_emitidas = set()
    for key, e in fams.items():
        out.append(f'0 {fid[key]} FAM')
        both = e['husb'] and e['wife']
        no_union = fam_flags[key]['no_union']
        if no_union:
            unsup.append((e['husb'], e['wife']))
            mar = '' if not e['mar'] else f"fecha={cv(e['mar'][0])}; lugar={cv(e['mar'][1])}"
            out += gedcom_text(1, 'NOTE', 'Union conyugal NO acreditada en 02_Relaciones (sin fila Conyuge/Pareja): esta familia solo agrupa a los padres documentados de sus hijos; NO modelar como matrimonio. Datos de matrimonio de la hoja 30: ' + (mar or 'ninguno'))
        for c in e['ch']:
            for ev in (evidencia or {}).get(c, []):
                if ev['progenitor'] not in (e['husb'], e['wife']):
                    continue
                priv_ev = priv_map.get(ev['hijo'], True) or priv_map.get(ev['progenitor'], True)
                ctx = ev['rel']
                if priv_ev:
                    lit = '[literal omitido por privacidad: hijo o progenitor privado (ver sidecar CSV de Genealogia)]'
                    stats['evidencia_literal_omitido_privacidad'] += 1
                else:
                    lit = '«' + scrub(ev['literal'], ctx + ':literal', omis) + '»'
                if ev.get('propio'):
                    lit += ' | texto propio de la fila: ' + ('[omitido por privacidad]' if priv_ev else '«' + scrub(ev['propio'], ctx + ':texto_propio', omis) + '»')
                if ev['origen']:
                    lit = f"igual que {ev['origen']} (misma acta, ese literal nombra a ambos progenitores): " + lit
                salv = scrub(ev['salvedades'], ctx + ':salvedades', omis) if not priv_ev else (f"Esta relación tiene salvedades de lectura; consulte sidecar {ev['rel']}" if ev['salvedades'] else 'ninguna declarada en el sidecar')
                salv = salv or 'ninguna declarada en el sidecar'
                txt = (f"EVIDENCIA RELACION {ev['rel']} | hijo {ev['hijo']} | {ev['tipo']} | progenitor {ev['progenitor']} | doc {ev['doc']} ({scrub(ev['tipo_doc'], ctx + ':tipo_doc', omis)}) | "
                       f"alcance de la validacion: {scrub(ev['alcance'], ctx + ':alcance', omis)} | literal leido: {lit} | salvedades: {salv} | "
                       f"citas candidatas de 02 (NO cotejadas con el literal, no son prueba de esta relacion): {ev['citas'] or 'ninguna'} | "
                       f"Nota de FAMILIA: se atribuye solo a la relacion {ev['rel']}; lectura incierta no es identidad confirmada.")
                out += gedcom_text(1, 'NOTE', nm(txt))
                stats['evidencia_bloques'] += 1; ev_emitidas.add(ev['rel'])
        if e.get('mar_conflict'):
            out += gedcom_text(1, 'NOTE', 'Datos de matrimonio CONTRADICTORIOS entre conyuges (no se escribe MARR): ' + ' || '.join(f'fecha={cv(a)}; lugar={cv(b)}' for a, b in e['mar_conflict']))
        if e['husb']:
            out.append(f"1 HUSB @{e['husb']}@")
        if e['wife']:
            out.append(f"1 WIFE @{e['wife']}@")
        for c in e['ch']:
            out.append(f'1 CHIL @{c}@')
        if e['mar'] and not no_union:
            out.append('1 MARR')
            g, note = parse_date(e['mar'][0], warn, f'{key} MARR')
            if g:
                out.append('2 DATE ' + g)
            if e['mar'][1]:
                out.append('2 PLAC ' + nm(e['mar'][1]))
            if note:
                out += gedcom_text(2, 'NOTE', 'Fecha original en el Excel: ' + note)
            for c, priv in fam_bound.get(key, []):
                out += cite_lines(c, 2, priv)
                stats['citas_mapeadas'] += 1; stats['citas_a_evento'] += 1
    for sid in sorted(used):
        s = S[sid]
        motivos = []
        if any(x['ID_Fuente'] == sid and nm(x['Derechos_privacidad']).lower() not in PUBLIC_OK for x in D):
            motivos.append('documento no publico')
        if sid in SRC_RESTRICT_DEFAULT:
            motivos.append('titulo con persona concreta o localizador de difuntos (decision por defecto, pendiente de validar)')
        if any(int(y) > CURRENT_YEAR_LIMIT for y in re.findall(r'\b(19\d\d|20\d\d)\b', nm(s['Titulo']) + ' ' + nm(s['Repositorio']))):
            motivos.append('titulo con fecha reciente (posible persona viva)')
        if SRC_PERSONAL.search(' '.join(nm(s[k]) for k in ('Tipo_fuente', 'Repositorio', 'URL', 'Autor_o_entidad'))):
            motivos.append('testimonio, mensaje o correo personal')
        if SRC_IDENTITY.search(nm(s['Titulo']) + ' ' + nm(s['Tipo_fuente']) + ' ' + nm(s['Autor_o_entidad'])):
            motivos.append('fuente de documento de identidad')
        if motivos:
            omis['fuente_restringida'].append(sid)
            omis['fuente_restringida_motivos'].append(f"{sid}: {', '.join(motivos)}")
            out += [f'0 @{sid}@ SOUR', f'1 TITL Fuente restringida {sid} (ver Excel)', '1 NOTE Titulo, autor, repositorio, signatura y URL omitidos por privacidad (ver Excel)']
            continue
        out += ['0 @%s@ SOUR' % sid] + gedcom_text(1, 'TITL', scrub(s['Titulo'], f'{sid}:Titulo', omis))
        if s['Autor_o_entidad']:
            if LOC_UNSAFE.search(nm(s['Autor_o_entidad'])) or FREE_IDENTITY.search(nm(s['Autor_o_entidad'])):
                omis['fuente_campo'].append(f'{sid}:Autor_o_entidad')
            else:
                out += gedcom_text(1, 'AUTH', nm(s['Autor_o_entidad']))
        campos = {'Repositorio': s['Repositorio'], 'Signatura': s['Signatura'], 'URL': s['URL']}
        for k, v in campos.items():
            if nm(v) and (LOC_UNSAFE.search(nm(v)) if k != 'URL' else (URL_UNSAFE.search(nm(v)) or url_private(nm(v)))):
                campos[k] = ''
                omis['fuente_campo'].append(f'{sid}:{k}')
            elif nm(v):
                campos[k] = scrub(v, f'{sid}:{k}', omis)
        out += gedcom_text(1, 'NOTE', nm(f"Tipo: {scrub(cv(s['Tipo_fuente']), sid + ':Tipo_fuente', omis)}; Original/derivada: {scrub(cv(s['Original_o_derivada']), sid + ':Original', omis)}; Calidad: {scrub(cv(s['Calidad_fuente']), sid + ':Calidad', omis)}; Repositorio: {cv(campos['Repositorio'])}; Signatura: {cv(campos['Signatura'])}; URL: {cv(campos['URL'])}"))
    for mid, did in sorted(media_used.items()):
        x = next(d for d in D if d['ID_Documento'] == did)
        rel, sha = media[did]
        out += [f'0 @{mid}@ OBJE', f'1 FILE {rel}', f"2 FORM {os.path.splitext(rel)[1].lstrip('.').lower() or 'bin'}", f"2 TITL {cv(x['Tipo_documento'])} ({did})"]
        out += gedcom_text(1, 'NOTE', f'SHA256 verificado contra el archivo: {sha}')
    out.append('0 TRLR')
    # Cruce con 02
    rep = []
    for key, e in fams.items():
        for c in e['ch']:
            pa = [r for r in R if r['ID_Persona_1'] == c and r['Tipo_relacion'] == 'Padre biológico']
            ma = [r for r in R if r['ID_Persona_1'] == c and r['Tipo_relacion'] == 'Madre biológica']
            if not ((not e['husb'] or any(r['ID_Persona_2'] == e['husb'] for r in pa)) and (not e['wife'] or any(r['ID_Persona_2'] == e['wife'] for r in ma))):
                rep.append(f"hijo {c}: familia ({e['husb'] or '-'},{e['wife'] or '-'}) NO respaldada en 02")
        if not (e['husb'] and e['wife']):
            rep.append(f"familia ({e['husb'] or '-'},{e['wife'] or '-'}): un solo progenitor en 30")
    for a, b in unsup:
        rep.append(f'pareja ({a},{b}): sin fila Conyuge en 02')
    hx = lambda v: bool(re.fullmatch(r'[0-9a-fA-F]{64}', str(v or '').strip()))
    fuera = [c for c in C if c['ID_Persona'] not in byid]
    vacia = sum(1 for c in fuera if not c['ID_Persona'])
    recon = stats['citas_mapeadas'] + stats['excluidas_confianza'] + stats['excluidas_sin_fuente'] + len(fuera)
    if recon != len(C):
        raise MapperError(f'reconciliacion de citas fallida: {recon} != {len(C)}')
    if evidencia:
        faltan = sorted(ev['rel'] for v in evidencia.values() for ev in v if ev['rel'] not in ev_emitidas)
        if faltan:
            raise MapperError('evidencia sin bloque (el progenitor no es HUSB/WIFE de la familia del hijo; incoherencia entre 02 y las personas): ' + ', '.join(faltan))
    summary = {
        'personas': len(P), 'familias': len(fams), 'fuentes_citadas': len(used),
        'citas_filas_total': len(C), 'citas_fuera_del_alcance': len(fuera), 'citas_fuera_persona_vacia': vacia,
        'personas_privadas_RESN': stats['personas_privadas'], 'evidencia_bloques': stats['evidencia_bloques'], 'evidencia_literal_omitido_privacidad': stats['evidencia_literal_omitido_privacidad'], 'docs_hash_no_valido_piloto': stats['docs_sin_hash_valido'], 'citas_de_personas_piloto': sum(len(v) for v in cit.values()),
        'citas_mapeadas': stats['citas_mapeadas'], 'citas_ligadas_a_evento': stats['citas_a_evento'], 'citas_ligadas_a_persona': stats['citas_a_persona'],
        'citas_persona_por_motivo': dict(reasons),
        'citas_excluidas_confianza': stats['excluidas_confianza'],
        'citas_excluidas_sin_fuente': stats['excluidas_sin_fuente'],
        'docs_total': len(D), 'docs_sha256_64hex': sum(hx(x['SHA256']) for x in D),
        'notas_personales_omitidas_privacidad': len(omis['nota_personal']),
        'citas_con_campos_omitidos_privacidad': len(omis['cita_campos']),
        'localizadores_omitidos_privacidad': len(omis['localizador']),
        'omisiones_privacidad_ids': {'nota_personal': sorted(omis['nota_personal']), 'localizador': sorted(omis['localizador']), 'cita_campos': sorted(omis['cita_campos']), 'fuente_restringida': sorted(omis['fuente_restringida']), 'fuente_restringida_motivos': sorted(omis['fuente_restringida_motivos']), 'fuente_campo': sorted(omis['fuente_campo']), 'libre_identidad': sorted(set(omis['libre_identidad'])), 'libre_enlace_privado': sorted(set(omis['libre_enlace_privado'])), 'libre_correo': sorted(set(omis['libre_correo'])), 'libre_numero_largo': sorted(set(omis['libre_numero_largo'])), 'documento_campos': sorted(set(omis['documento_campos']))},
        'documentos_con_campos_omitidos_privacidad': len(set(omis['documento_campos'])),
        'fuentes_restringidas_privacidad': len(omis['fuente_restringida']), 'campos_fuente_omitidos_privacidad': len(omis['fuente_campo']),
        'media_manifiesto': (len(media) if media_manifest else 'no usado'), 'media_objetos_unicos': len(media_used), 'media_objetos_enlazados': stats['media_objetos_enlazados'], 'media_no_migrados_por_motivo': dict(media_reasons),
        'docs_persona_piloto': sum(1 for x in D if x['ID_Persona_principal'] in byid),
    }
    # Publicacion atomica: solo tras validar y reconciliar todo; el archivo anterior se conserva ante cualquier error.
    out_dir = os.path.dirname(out_path) or '.'
    os.makedirs(out_dir, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix='.mapper-', suffix='.tmp', dir=out_dir)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as fh:
            fh.write('\n'.join(out) + '\n')
            fh.flush(); os.fsync(fh.fileno())
        os.replace(tmp, out_path)
    except BaseException:
        try:
            if os.path.exists(tmp):
                os.unlink(tmp)
        except OSError:
            # no enmascarar el error original; avisar del residuo (puede contener datos privados, permisos 0600) sin revelar su contenido
            print(f'AVISO: no se pudo borrar el temporal privado {tmp}; borrarlo manualmente', file=sys.stderr)
        raise
    return summary, warn, rep, unsup


def main(argv):
    manifest = None
    args = []
    for a in argv[1:]:
        if a.startswith('--media-manifest='):
            if manifest is not None or not a.split('=', 1)[1]:
                print('uso: python3 mapper.py entrada.xlsx [salida.ged] [--media-manifest=RUTA.csv]'); return 2
            manifest = a.split('=', 1)[1]
        elif a.startswith('-'):
            print('uso: python3 mapper.py entrada.xlsx [salida.ged] [--media-manifest=RUTA.csv]'); return 2
        else:
            args.append(a)
    if not 1 <= len(args) <= 2:
        print('uso: python3 mapper.py entrada.xlsx [salida.ged] [--media-manifest=RUTA.csv]'); return 2
    if os.path.isfile(args[0]):
        _msg = check_output_path(args[0], args[1] if len(args) > 1 else 'out/salida.ged')
        if _msg:
            print('ERROR:', _msg); return 2
    try:
        summary, warn, rep, unsup = run(args[0], args[1] if len(args) > 1 else 'out/salida.ged', media_manifest=manifest)
    except MapperError as ex:
        print('ERROR (fail closed):', ex); return 1
    for k, v in summary.items():
        print(f'{k}: {v}')
    print('advertencias:', len(warn)); [print('  ', w) for w in warn]
    print('cruce con 02:'); [print('  ', r) for r in rep] if rep else print('   sin discrepancias')
    print('familias con union no acreditada (sin MARR):', unsup)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
